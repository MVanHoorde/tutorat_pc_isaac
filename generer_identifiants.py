#!/usr/bin/env python3
"""
generer_identifiants.py — Crée les identifiants Résonance à partir des trombinoscopes PDF.

Pour chaque élève trouvé dans les trombinoscopes, le script tire au hasard un
identifiant (PC-XXXXX) et un mot de passe, puis écrit dans le dossier de sortie :

  correspondance_resonance.xlsx   registre noms ↔ identifiants, cumulatif.
                                  À conserver dans OneDrive, réservé aux professeurs.
  coupons_AAAAMMJJ-HHMM.pdf       coupons à imprimer et découper (nouveaux élèves seulement).
                                  Contient noms et mots de passe : à supprimer après impression.
  comptes_serveur_AAAAMMJJ-HHMM.csv
                                  comptes à charger sur le serveur : identifiant, classe,
                                  empreinte du mot de passe. AUCUN NOM.

Relancer le script avec le même dossier de sortie ne change pas les identifiants
déjà attribués : seuls les nouveaux élèves reçoivent un compte et un coupon.

    python generer_identifiants.py trombi_1re2.pdf trombi_tle1.pdf --sortie "C:/Users/.../OneDrive/Résonance"
    python generer_identifiants.py trombi.pdf --simulation          # lit, n'écrit rien
    python generer_identifiants.py trombi.pdf --classe "1re spé 2"  # impose le nom de classe

Par défaut, les noms ne s'affichent pas dans le terminal (masqués comme dans
analyser_trombi.py) ; --afficher-noms pour les voir.

Dépendances : pip install pymupdf openpyxl
"""

import argparse
import base64
import csv
import hashlib
import secrets
import sys
import unicodedata
from datetime import datetime
from html import escape
from pathlib import Path

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        sys.exit("Installez PyMuPDF :  pip install pymupdf")
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    sys.exit("Installez openpyxl :  pip install openpyxl")


REGISTRE = "correspondance_resonance.xlsx"
COLONNES = ["Identifiant", "Classe", "Nom", "Prénom", "Ajouté le", "Source"]
SITE = "mvanhoorde.github.io/tutorat_pc_isaac"

# Sans 0/O, 1/I/L : rien qui se confonde à la lecture d'un coupon.
ALPHABET_ID = "ABCDEFGHJKMNPQRSTUVWXYZ23456789"
MOTS = """abeille acier agate aimant algue ambre ancre arbre argent atome aube aurore
avion azur baleine bambou banjo biscuit bison boussole braise brume cactus camion
canard canyon carbone castor cerise chamois chaton cigale citron cobalt colibri comete
corail cristal cyclone dauphin delta diamant domino dune eclair ecureuil electron
etoile falaise faucon fenouil figue flocon fougere fusee galaxie gazelle geyser girafe
glacier granit grenat grillon hamac hibou horizon iceberg iris jade jaguar kiwi koala
lagune lama laser lavande lezard lichen lotus lynx magnet mammouth mangue marmotte
meteore mimosa mistral nectar neige nuage oasis ocean olive onde opale orage orbite
orchidee ours panda papaya pepite phare photon pingouin pixel planete plasma pollen
poulpe prisme puma quartz radar radis renard requin rubis safran saphir saturne
sequoia silex sirocco soleil sonar spirale tamis tigre titane topaze tornade toucan
tulipe turbo vanille vapeur volcan zebre zenith""".split()


# --------------------------------------------------------------------------
# Lecture des trombinoscopes
# --------------------------------------------------------------------------

def masquer(texte):
    return "".join("9" if c.isdigit() else "A" if c.isupper() else "a" if c.islower() else c
                   for c in texte)


def blocs_texte(page):
    blocs = []
    for bloc in page.get_text("dict")["blocks"]:
        if bloc.get("type") != 0:
            continue
        spans = [s for l in bloc["lines"] for s in l["spans"] if s["text"].strip()]
        if not spans:
            continue
        x0, y0, x1, y1 = bloc["bbox"]
        blocs.append({"x0": x0, "y0": y0, "x1": x1, "y1": y1,
                      "texte": " ".join(s["text"].strip() for s in spans),
                      "taille": max(s["size"] for s in spans)})
    return blocs


def photos(page):
    """Rectangles des photos, en écartant ce qui n'a pas la taille dominante (logo...)."""
    rects = []
    for img in page.get_images(full=True):
        try:
            rects += page.get_image_rects(img[0])
        except Exception:
            pass
    uniques = {(round(r.x0), round(r.y0)): r for r in rects}.values()
    if not uniques:
        return []
    tailles = {}
    for r in uniques:
        tailles.setdefault((round(r.width), round(r.height)), []).append(r)
    dominante = max(tailles.values(), key=len)
    return sorted(dominante, key=lambda r: (round(r.y0), r.x0))


def classe_en_tete(blocs):
    """Le plus gros texte en haut de page ; « Classe : Première 2 » → « Première 2 »."""
    hauts = [b for b in blocs if b["y0"] < 90]
    if not hauts:
        return None
    titre = max(hauts, key=lambda b: b["taille"])["texte"]
    return titre.split(":", 1)[1].strip() if ":" in titre else titre.strip()


def separer_nom_prenom(texte):
    """« DUPONT MARTIN Jean-Luc » → ("DUPONT MARTIN", "Jean-Luc")."""
    mots = texte.split()
    i = 0
    while i < len(mots) and any(c.isalpha() for c in mots[i]) and mots[i] == mots[i].upper():
        i += 1
    return " ".join(mots[:i]), " ".join(mots[i:])


def lire_trombi(chemin, classe_imposee=None):
    """Retourne (classe, élèves, alertes). Un élève = {nom, prenom, page}."""
    doc = fitz.open(chemin)
    classe, eleves, alertes = classe_imposee, [], []
    for num, page in enumerate(doc, start=1):
        blocs = blocs_texte(page)
        if classe is None:
            classe = classe_en_tete(blocs)
        utilises = set()
        for r in photos(page):
            cx = (r.x0 + r.x1) / 2
            dessous = [b for b in blocs
                       if 0 <= b["y0"] - r.y1 <= 40
                       and abs((b["x0"] + b["x1"]) / 2 - cx) <= r.width * 0.6]
            dessous.sort(key=lambda b: b["y0"])
            if not dessous:
                alertes.append(f"page {num} : photo sans nom en ({r.x0:.0f}, {r.y0:.0f})")
                continue
            utilises.update(id(b) for b in dessous)
            texte = " ".join(b["texte"] for b in dessous)
            nom, prenom = separer_nom_prenom(texte)
            if not nom or not prenom:
                alertes.append(f"page {num} : nom ou prénom introuvable dans « {texte} »")
            eleves.append({"nom": nom or texte, "prenom": prenom, "page": num})
        # Noms de même taille que les autres mais sous aucune photo.
        tailles_noms = {round(b["taille"]) for b in blocs if id(b) in utilises}
        for b in blocs:
            if id(b) not in utilises and round(b["taille"]) in tailles_noms and 90 < b["y0"] < page.rect.height - 50:
                alertes.append(f"page {num} : texte sans photo « {b['texte']} »")
    doc.close()
    if not classe:
        alertes.append("classe introuvable dans l'en-tête : utilisez --classe")
        classe = Path(chemin).stem
    return classe, eleves, alertes


# --------------------------------------------------------------------------
# Identifiants et mots de passe
# --------------------------------------------------------------------------

def cle(classe, nom, prenom):
    norm = lambda s: unicodedata.normalize("NFC", " ".join(str(s or "").split())).casefold()
    return (norm(classe), norm(nom), norm(prenom))


def nouvel_identifiant(pris):
    while True:
        ident = "PC-" + "".join(secrets.choice(ALPHABET_ID) for _ in range(5))
        if ident not in pris:
            pris.add(ident)
            return ident


def nouveau_mot_de_passe():
    a, b = secrets.choice(MOTS), secrets.choice(MOTS)
    return f"{a.capitalize()}-{b.capitalize()}-{secrets.randbelow(90) + 10}"


def empreinte(mot_de_passe, iterations=600_000):
    """PBKDF2-SHA256, au format lisible par Django et la plupart des frameworks."""
    sel = secrets.token_hex(12)
    h = hashlib.pbkdf2_hmac("sha256", mot_de_passe.encode(), sel.encode(), iterations)
    return f"pbkdf2_sha256${iterations}${sel}${base64.b64encode(h).decode()}"


# --------------------------------------------------------------------------
# Écriture
# --------------------------------------------------------------------------

def lire_registre(chemin):
    if not chemin.exists():
        return []
    ws = load_workbook(chemin).active
    lignes = list(ws.iter_rows(values_only=True))
    entete = list(lignes[0])
    return [dict(zip(entete, l)) for l in lignes[1:] if l and l[0]]


def ecrire_registre(chemin, lignes):
    wb = Workbook()
    ws = wb.active
    ws.title = "Correspondance"
    ws.append(COLONNES)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="151B2B")
    for l in sorted(lignes, key=lambda l: (str(l["Classe"]), str(l["Nom"]), str(l["Prénom"]))):
        ws.append([l.get(c) for c in COLONNES])
    for col, largeur in zip("ABCDEF", (13, 16, 24, 20, 17, 30)):
        ws.column_dimensions[col].width = largeur
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(chemin)
    except PermissionError:
        sys.exit(f"\n  ⚠  Impossible d'écrire {chemin.name} : le fichier est ouvert (Excel, aperçu OneDrive…).\n"
                 f"     Fermez-le puis relancez la même commande. Rien n'a été créé.\n")


CSS_COUPON = """
* { font-family: sans-serif; }
.titre { font-size: 8pt; color: #5C6679; margin: 0; }
.nom { font-size: 12pt; font-weight: bold; margin: 3pt 0 0 0; }
.classe { font-size: 9pt; color: #5C6679; margin: 0 0 6pt 0; }
.lab { font-size: 7.5pt; color: #5C6679; margin: 0; }
.val { font-family: monospace; font-size: 13pt; font-weight: bold; margin: 0 0 3pt 0; }
.pied { font-size: 7pt; color: #5C6679; margin: 4pt 0 0 0; }
"""


def ecrire_coupons(chemin, comptes):
    """10 coupons par page A4 (2 × 5), avec traits de coupe."""
    doc = fitz.open()
    L, H, marge = 595.3, 841.9, 28
    lc, hc = (L - 2 * marge) / 2, (H - 2 * marge) / 5
    for i, c in enumerate(comptes):
        if i % 10 == 0:
            page = doc.new_page(width=L, height=H)
            for k in range(1, 5):
                y = marge + k * hc
                page.draw_line((marge, y), (L - marge, y), color=(.6, .6, .6), dashes="[3 3] 0", width=.5)
            page.draw_line((L / 2, marge), (L / 2, H - marge), color=(.6, .6, .6), dashes="[3 3] 0", width=.5)
            page.draw_rect(fitz.Rect(marge, marge, L - marge, H - marge), color=(.6, .6, .6), dashes="[3 3] 0", width=.5)
        col, lig = (i % 10) % 2, (i % 10) // 2
        x0, y0 = marge + col * lc, marge + lig * hc
        page.draw_rect(fitz.Rect(x0 + 14, y0 + 16, x0 + 17, y0 + hc - 16), color=None, fill=(1, .37, .64))
        html = (
            f'<p class="titre">RÉSONANCE — tutorat en physique-chimie</p>'
            f'<p class="nom">{escape(c["Prénom"])} {escape(c["Nom"])}</p>'
            f'<p class="classe">{escape(str(c["Classe"]))}</p>'
            f'<p class="lab">Identifiant</p><p class="val">{escape(c["Identifiant"])}</p>'
            f'<p class="lab">Mot de passe</p><p class="val">{escape(c["mdp"])}</p>'
            f'<p class="pied">{SITE} · Garde ce coupon, ne le partage pas.</p>'
        )
        page.insert_htmlbox(fitz.Rect(x0 + 26, y0 + 12, x0 + lc - 12, y0 + hc - 8), html, css=CSS_COUPON)
    doc.save(chemin)
    doc.close()


def ecrire_serveur(chemin, comptes):
    with open(chemin, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["identifiant", "classe", "empreinte_mot_de_passe"])
        for c in comptes:
            w.writerow([c["Identifiant"], c["Classe"], empreinte(c["mdp"])])


# --------------------------------------------------------------------------

def main():
    # Console Windows : sans ça, les accents et « ⚠ » peuvent faire planter l'affichage.
    for flux in (sys.stdout, sys.stderr):
        try:
            flux.reconfigure(encoding="utf-8")
        except AttributeError:
            pass
    ap = argparse.ArgumentParser(description="Crée les identifiants Résonance à partir des trombinoscopes PDF.")
    ap.add_argument("pdf", nargs="+", help="trombinoscope(s) PDF, un par classe")
    ap.add_argument("--sortie", default="sortie_resonance", help="dossier de sortie (idéalement dans OneDrive)")
    ap.add_argument("--classe", help="impose le nom de classe (si un seul PDF)")
    ap.add_argument("--simulation", action="store_true", help="lit les PDF et affiche le bilan, n'écrit rien")
    ap.add_argument("--afficher-noms", action="store_true", help="affiche les vrais noms dans le terminal")
    args = ap.parse_args()
    if args.classe and len(args.pdf) > 1:
        sys.exit("--classe ne s'utilise qu'avec un seul PDF.")

    voir = (lambda s: s) if args.afficher_noms else masquer
    sortie = Path(args.sortie)
    registre = lire_registre(sortie / REGISTRE)
    connus = {cle(l["Classe"], l["Nom"], l["Prénom"]) for l in registre}
    pris = {l["Identifiant"] for l in registre}
    maintenant = datetime.now()
    nouveaux, total_alertes = [], 0

    print()
    for chemin in args.pdf:
        classe, eleves, alertes = lire_trombi(chemin, args.classe)
        deja = 0
        doublons = set()
        for e in eleves:
            k = cle(classe, e["nom"], e["prenom"])
            if k in connus:
                deja += 1
                continue
            if k in doublons:
                alertes.append(f"doublon dans le PDF : « {voir(e['nom'] + ' ' + e['prenom'])} »")
                continue
            doublons.add(k)
            nouveaux.append({"Identifiant": nouvel_identifiant(pris), "Classe": classe,
                             "Nom": e["nom"], "Prénom": e["prenom"],
                             "Ajouté le": maintenant.strftime("%d/%m/%Y %H:%M"),
                             "Source": Path(chemin).name, "mdp": nouveau_mot_de_passe()})
        print(f"  {Path(chemin).name}")
        print(f"      classe          : {classe}")
        print(f"      élèves lus      : {len(eleves)}")
        print(f"      déjà inscrits   : {deja}")
        print(f"      nouveaux        : {len(doublons)}")
        for a in alertes:
            # Les alertes peuvent citer un nom : masqué sauf --afficher-noms.
            if "«" in a:
                debut, reste = a.split("«", 1)
                a = debut + "« " + voir(reste.rsplit("»", 1)[0].strip()) + " »"
            print(f"      ⚠  {a}")
        total_alertes += len(alertes)
        print()

    if args.afficher_noms and nouveaux:
        print("  Nouveaux comptes :")
        for c in nouveaux:
            print(f"      {c['Identifiant']}  {c['Classe']:<14} {c['Nom']} {c['Prénom']}")
        print()

    if args.simulation:
        print("  Simulation : aucun fichier écrit.\n")
        return
    if not nouveaux:
        print("  Aucun nouvel élève : rien à écrire.\n")
        return

    sortie.mkdir(parents=True, exist_ok=True)
    horodatage = maintenant.strftime("%Y%m%d-%H%M")
    ecrire_registre(sortie / REGISTRE, registre + nouveaux)
    ecrire_coupons(sortie / f"coupons_{horodatage}.pdf", nouveaux)
    ecrire_serveur(sortie / f"comptes_serveur_{horodatage}.csv", nouveaux)

    print(f"  {len(nouveaux)} compte(s) créé(s) dans {sortie.resolve()}")
    print(f"      {REGISTRE:<36} registre complet ({len(registre) + len(nouveaux)} élèves) — OneDrive, profs uniquement")
    print(f"      coupons_{horodatage}.pdf{'':<11} à imprimer, puis à SUPPRIMER")
    print(f"      comptes_serveur_{horodatage}.csv{'':<4} pour le serveur — sans aucun nom")
    if total_alertes:
        print(f"\n  ⚠  {total_alertes} alerte(s) : vérifiez ces élèves dans le registre avant d'imprimer.")
    print()


if __name__ == "__main__":
    main()
